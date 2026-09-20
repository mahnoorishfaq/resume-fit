"""
The analysis engine.

Three jobs:
  1. pull text out of a resume file
  2. work out which skills it contains and how well it fits a job
  3. check whether the resume itself is well built (the ATS checks)

No large models are downloaded. Everything here runs in under a second,
which matters when the app is on a free hosting tier.
"""

import io
import re
from collections import Counter

from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

import skills_db

# ----------------------------------------------------------------------
# 1. READING THE FILE
# ----------------------------------------------------------------------

def extract_text(uploaded_file, filename=None):
    """
    Pull plain text out of a PDF, DOCX or TXT upload.
    Returns (text, error_message). error_message is None on success.

    Works with both Flask's FileStorage (which exposes .filename) and
    Streamlit's UploadedFile (which exposes .name).
    """
    name = (filename
            or getattr(uploaded_file, "filename", None)
            or getattr(uploaded_file, "name", "")).lower()
    raw = uploaded_file.read()

    if name.endswith(".pdf"):
        return _read_pdf(raw)
    if name.endswith(".docx"):
        return _read_docx(raw)
    if name.endswith(".txt"):
        try:
            return raw.decode("utf-8", errors="ignore"), None
        except Exception as e:
            return "", f"Could not read the text file: {e}"

    return "", "Upload a PDF, DOCX or TXT file."


def _read_pdf(raw):
    try:
        import pdfplumber
        text_parts = []
        with pdfplumber.open(io.BytesIO(raw)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        text = "\n".join(text_parts)
        if not text.strip():
            return "", ("No text found. This PDF is probably a scan or an "
                        "image. Export your resume as a text-based PDF.")
        return text, None
    except Exception as e:
        return "", f"Could not read the PDF: {e}"


def _read_docx(raw):
    try:
        import docx
        document = docx.Document(io.BytesIO(raw))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        # Tables often hold skills sections
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)
        return "\n".join(parts), None
    except Exception as e:
        return "", f"Could not read the Word file: {e}"


# ----------------------------------------------------------------------
# 2. FINDING SKILLS
# ----------------------------------------------------------------------

def find_skills(text):
    """
    Return the set of known skills present in the text.
    Matches on whole words so 'r' doesn't match every word containing r.
    """
    if not text:
        return set()

    lowered = " " + re.sub(r"[^a-z0-9+#/.\- ]", " ", text.lower()) + " "
    lowered = re.sub(r"\s+", " ", lowered)

    found = set()
    for skill, meta in skills_db.SKILLS.items():
        for alias in meta["aliases"]:
            alias = alias.strip()
            if not alias:
                continue
            # Escape regex characters in things like "c++" or "scikit-learn"
            pattern = r"(?<![a-z0-9])" + re.escape(alias) + r"(?![a-z0-9])"
            if re.search(pattern, lowered):
                found.add(skill)
                break
    return found


# ----------------------------------------------------------------------
# 3. SCORING THE FIT
# ----------------------------------------------------------------------

def text_similarity(resume_text, job_text):
    """
    Overall wording overlap between resume and job description, 0-100.
    TF-IDF cosine similarity: fast, no model download, good enough here.
    """
    if not resume_text.strip() or not job_text.strip():
        return 0.0
    try:
        vectorizer = TfidfVectorizer(stop_words="english", ngram_range=(1, 2))
        matrix = vectorizer.fit_transform([resume_text, job_text])
        score = cosine_similarity(matrix[0:1], matrix[1:2])[0][0]
        return round(float(score) * 100, 1)
    except ValueError:
        return 0.0


def analyse_fit(resume_text, job_text):
    """
    Compare a resume against one job description.
    Returns a dict the UI can render directly.
    """
    resume_skills = find_skills(resume_text)
    job_skills = find_skills(job_text)

    matched = sorted(resume_skills & job_skills)
    missing = sorted(job_skills - resume_skills)
    extra = sorted(resume_skills - job_skills)

    # Skill coverage is the signal that actually matters
    if job_skills:
        coverage = len(matched) / len(job_skills) * 100
    else:
        coverage = 0.0

    wording = text_similarity(resume_text, job_text)

    # Weight coverage far higher than wording overlap. Matching the
    # required skills matters more than using the same vocabulary.
    score = round(coverage * 0.75 + wording * 0.25)
    score = max(0, min(100, score))

    return {
        "score": score,
        "coverage": round(coverage),
        "wording": round(wording),
        "matched": matched,
        "missing": missing,
        "extra": extra,
        "verdict": verdict_for(score),
    }


def verdict_for(score):
    if score >= 70:
        return ("Strong match",
                "Your resume covers most of what this role asks for. "
                "Worth applying.")
    if score >= 50:
        return ("Worth applying",
                "You cover the core of this role. Closing one or two gaps "
                "below would make it much stronger.")
    if score >= 30:
        return ("Stretch role",
                "There's real distance here. Apply if you want it, but "
                "work through the gaps first.")
    return ("Not yet",
            "This role is asking for a different profile. The skills below "
            "are the shortest route toward it.")


def learning_plan(missing_skills, limit=6):
    """
    Turn missing skills into an ordered list of things to learn.
    High-leverage skills come first because they appear in many job ads
    and rarely in student resumes.
    """
    plan = []
    for skill in missing_skills:
        plan.append({
            "skill": skill,
            "category": skills_db.category_of(skill),
            "note": skills_db.learning_note(skill),
            "priority": skill in skills_db.HIGH_LEVERAGE,
        })
    plan.sort(key=lambda item: (not item["priority"], item["skill"]))
    return plan[:limit]


def recommend_jobs(resume_text, jobs, top_n=5):
    """Rank every job in the library against this resume."""
    results = []
    for job in jobs:
        blob = f"{job['title']} {job.get('company','')} {job['description']}"
        analysis = analyse_fit(resume_text, blob)
        results.append({**job, **analysis})
    results.sort(key=lambda j: j["score"], reverse=True)
    return results[:top_n]


# ----------------------------------------------------------------------
# 4. RESUME HEALTH CHECKS
# ----------------------------------------------------------------------

ACTION_VERBS = {
    "built", "designed", "developed", "led", "created", "implemented",
    "improved", "reduced", "increased", "automated", "analysed", "analyzed",
    "deployed", "launched", "managed", "trained", "optimised", "optimized",
    "delivered", "researched", "engineered", "shipped",
}

WEAK_PHRASES = [
    "responsible for", "duties included", "worked on", "helped with",
    "team player", "hard working", "hard-working", "go-getter",
    "think outside the box", "detail oriented", "detail-oriented",
]


def health_checks(text, filename=""):
    """
    Heuristic checks on the resume itself, independent of any job.
    Each returns: name, passed (True/False/None), message.
    None means 'worth a look' rather than pass or fail.
    """
    checks = []
    lowered = text.lower()
    words = text.split()
    word_count = len(words)

    # Contact details
    has_email = bool(re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", text))
    has_phone = bool(re.search(r"(\+?\d[\d\s\-()]{8,}\d)", text))
    if has_email and has_phone:
        checks.append(("Contact details", True,
                       "Email and phone number both found."))
    elif has_email:
        checks.append(("Contact details", None,
                       "Email found, no phone number. Add one."))
    else:
        checks.append(("Contact details", False,
                       "No email address found. This alone can sink an "
                       "application."))

    # Links
    has_link = bool(re.search(r"(github\.com|linkedin\.com|gitlab\.com)", lowered))
    checks.append(("Portfolio links", has_link,
                   "GitHub or LinkedIn link found."
                   if has_link else
                   "No GitHub or LinkedIn link. For a technical role this is "
                   "expected."))

    # Length
    if word_count < 200:
        checks.append(("Length", False,
                       f"Only {word_count} words. This reads as thin - "
                       f"aim for 400 to 800."))
    elif word_count > 1200:
        checks.append(("Length", None,
                       f"{word_count} words is long for an early-career "
                       f"resume. Cut to the strongest material."))
    else:
        checks.append(("Length", True,
                       f"{word_count} words, which sits in a good range."))

    # Standard sections
    expected = ["experience", "education", "skill", "project"]
    present = [s for s in expected if s in lowered]
    if len(present) >= 3:
        checks.append(("Section headings", True,
                       f"Found {len(present)} standard sections. Parsers "
                       f"rely on these."))
    else:
        missing = [s for s in expected if s not in present]
        checks.append(("Section headings", False,
                       f"Missing clear headings for: {', '.join(missing)}. "
                       f"Automated parsers look for these words."))

    # Quantified results
    numbers = re.findall(r"\b\d+(?:\.\d+)?%?\b", text)
    meaningful = [n for n in numbers if len(n) > 1]
    if len(meaningful) >= 5:
        checks.append(("Quantified results", True,
                       f"{len(meaningful)} figures found. Numbers make "
                       f"claims concrete."))
    else:
        checks.append(("Quantified results", False,
                       "Few numbers found. 'Cut processing time by 40%' "
                       "beats 'improved processing time'."))

    # Action verbs
    starts = re.findall(r"(?:^|\n)\s*[-•*\u2022]?\s*(\w+)", text)
    verb_hits = sum(1 for w in starts if w.lower() in ACTION_VERBS)
    if verb_hits >= 4:
        checks.append(("Action verbs", True,
                       f"{verb_hits} bullet points open with a strong verb."))
    else:
        checks.append(("Action verbs", False,
                       "Start bullet points with what you did: built, "
                       "designed, reduced, automated."))

    # Filler phrases
    found_weak = [p for p in WEAK_PHRASES if p in lowered]
    if found_weak:
        checks.append(("Filler phrases", False,
                       f"Found: {', '.join(found_weak[:3])}. These take up "
                       f"space without saying anything."))
    else:
        checks.append(("Filler phrases", True,
                       "No common filler phrases found."))

    # File format
    if filename.lower().endswith(".pdf"):
        checks.append(("File format", True,
                       "PDF keeps your formatting intact across systems."))
    elif filename.lower().endswith(".docx"):
        checks.append(("File format", None,
                       "DOCX is accepted, but PDF is safer unless the "
                       "employer asks otherwise."))
    else:
        checks.append(("File format", None,
                       "Send PDF when you apply."))

    return checks


def health_score(checks):
    """Turn the checks into a single 0-100 number."""
    if not checks:
        return 0
    points = 0
    for _, passed, _ in checks:
        if passed is True:
            points += 1
        elif passed is None:
            points += 0.5
    return round(points / len(checks) * 100)
